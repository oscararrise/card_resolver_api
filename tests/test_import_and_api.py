import base64
import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from app import main
from app.database import Base, Credential


def test_sample_roster_end_to_end(tmp_path, monkeypatch):
    import os
    from pathlib import Path

    sample = Path(__file__).resolve().parents[2] / "upload" / "BASE NUEVA(BASE GENERAL)(in).csv"
    if not sample.exists():
        import pytest
        pytest.skip("private source CSV is not shipped with repository")
    monkeypatch.setenv("ADMIN_USER", "admin")
    monkeypatch.setenv("ADMIN_PASSWORD", "very-long-test-password")
    monkeypatch.setenv("SERVICE_TOKEN", "service-test-token")
    main.engine.dispose()
    from app.database import make_session
    main.engine, main.Session = make_session(f"sqlite:///{tmp_path}/test.db")
    auth = "Basic " + base64.b64encode(b"admin:very-long-test-password").decode()
    with TestClient(main.app) as client:
        response = client.post("/admin/imports", files={"file": (sample.name, sample.read_bytes(), "text/csv")},
                               headers={"Authorization": auth})
        assert response.status_code == 201, response.text
        batch = response.json()
        assert batch["cards"] > 500
        with main.Session() as session:
            result = session.scalars(select(Credential).where(Credential.batch_id == batch["id"],
                Credential.facility_code == 1761, Credential.card_number == 167794)).all()
            assert len(result) == 1 and result[0].eid == "48251"
        report = client.get(f'/admin/imports/{batch["id"]}/report', headers={"Authorization": auth})
        assert report.status_code == 200
        wb = load_workbook(io.BytesIO(report.content), read_only=True)
        assert {"SUMMARY", "LOADED", "ERRORS", "WARNINGS", "DUPLICATES"} <= set(wb.sheetnames)
        wb.close()
        assert client.post("/api/v1/cards/resolve", json={"raw": "89910480577"},
                           headers={"Authorization": "Bearer service-test-token"}).status_code == 503
        activation = client.post(f'/admin/imports/{batch["id"]}/activate', data={"reviewed": "true"},
                                 headers={"Authorization": auth})
        assert activation.status_code == 200, activation.text
        resolved = client.post("/api/v1/cards/resolve", json={"raw": "89910480577"},
                               headers={"Authorization": "Bearer service-test-token"})
        assert resolved.status_code == 200
        assert resolved.json()["employee"]["hibob_id"] == "48251"
        assert client.post("/api/v1/cards/decode", json={"raw": "89910480577"}).status_code == 401

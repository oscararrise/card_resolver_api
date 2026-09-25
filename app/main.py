import io
import os
import secrets
from urllib.parse import urlsplit

from fastapi import Depends, FastAPI, File, Form, Header, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.security import HTTPBasic, HTTPBasicCredentials
from openpyxl import Workbook
from pydantic import BaseModel, Field
from sqlalchemy import func, select

from .database import Active, Base, Batch, Credential, Issue, make_session
from .decoder import InvalidCard, decode
from .importer import InvalidUpload, stage
from .config import setting

engine, Session = make_session()
app = FastAPI(title="Card Resolver", docs_url=None, redoc_url=None)
basic = HTTPBasic(auto_error=False)


@app.on_event("startup")
def create_tables():
    Base.metadata.create_all(engine)


def db():
    with Session() as session:
        yield session


def api_auth(authorization: str | None = Header(default=None)):
    token = setting("SERVICE_TOKEN")
    if not token:
        raise HTTPException(503, "service token is not configured")
    if not authorization or not secrets.compare_digest(authorization, "Bearer " + token):
        raise HTTPException(401, "unauthorized", headers={"WWW-Authenticate": "Bearer"})


def admin_auth(credentials: HTTPBasicCredentials | None = Depends(basic)):
    user, password = os.getenv("ADMIN_USER", ""), setting("ADMIN_PASSWORD")
    if not user or not password:
        raise HTTPException(503, "admin authentication is not configured")
    if not credentials or not (secrets.compare_digest(credentials.username, user) and
                               secrets.compare_digest(credentials.password, password)):
        raise HTTPException(401, "unauthorized", headers={"WWW-Authenticate": 'Basic realm="Card imports"'})


def same_origin(request: Request):
    origin = request.headers.get("origin")
    if origin and urlsplit(origin).netloc.lower() != request.headers.get("host", "").lower():
        raise HTTPException(403, "cross-origin admin request blocked")


class CardRequest(BaseModel):
    raw: str = Field(min_length=1, max_length=20)
    format: str = "auto"


def parsed(data):
    try:
        return decode(data.raw, data.format)
    except InvalidCard as exc:
        raise HTTPException(422, str(exc)) from exc


@app.get("/health")
def health():
    return {"status": "ok"}


@app.post("/api/v1/cards/decode", dependencies=[Depends(api_auth)])
def decode_card(data: CardRequest):
    return {"raw": data.raw, "decoder": parsed(data)}


@app.post("/api/v1/cards/resolve", dependencies=[Depends(api_auth)])
def resolve_card(data: CardRequest, session=Depends(db)):
    result = parsed(data)
    active = session.get(Active, 1)
    if not active:
        raise HTTPException(503, "no active dataset")
    rows = session.scalars(select(Credential).where(
        Credential.batch_id == active.batch_id,
        Credential.facility_code == result["facility_code"],
        Credential.card_number == result["card_number"]).limit(2)).all()
    if len(rows) > 1:
        raise HTTPException(409, "ambiguous card")
    if not rows:
        raise HTTPException(404, "card not found")
    match = rows[0]
    return {"raw": data.raw, "decoder": result,
            "employee": {"hibob_id": match.eid, "name": match.full_name},
            "dataset_id": active.batch_id}


def batch_data(batch):
    return {"id": batch.id, "filename": batch.filename, "sha256": batch.sha256,
            "status": batch.status, "uploaded_at": batch.uploaded_at.isoformat(),
            "rows_read": batch.rows_read, "loaded": batch.loaded, "skipped": batch.skipped,
            "errors": batch.errors, "warnings": batch.warnings, "cards": batch.cards_count}


@app.get("/admin", response_class=HTMLResponse, dependencies=[Depends(admin_auth)])
def admin_page():
    return """<!doctype html><html lang="es"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Card Resolver · Importaciones</title><style>
body{font:16px system-ui;background:#f3f6f5;color:#152c28;max-width:850px;margin:3rem auto;padding:0 1rem}
main{background:#fff;padding:2rem;border-radius:14px;box-shadow:0 8px 30px #103f2d13}
button,input{font:inherit;padding:.65rem;margin:.35rem}button{background:#126a54;color:white;border:0;border-radius:7px;cursor:pointer}
button:disabled{opacity:.5}pre{white-space:pre-wrap;word-break:break-word;background:#f3f6f5;padding:1rem}
a{color:#126a54}li{margin:.75rem 0}</style><main><h1>Importaciones de tarjetas</h1>
<p>Sube el roster CSV o XLSX, revisa el reporte y activa la versión. Las reposiciones quedan pendientes de confirmar.</p>
<form id="upload"><input type="file" name="file" accept=".csv,.xlsx" required><button>Cargar y validar</button></form>
<pre id="result" role="status"></pre><h2>Versiones</h2><ul id="versions"></ul></main>
<script>
const result=document.querySelector('#result');
const adminBase=location.pathname.endsWith('/') ? location.pathname : location.pathname+'/';
async function refresh(){const r=await fetch(adminBase+'imports');const list=await r.json();
document.querySelector('#versions').replaceChildren(...list.map(b=>{
const li=document.createElement('li');li.append(document.createTextNode(`Versión ${b.id}: ${b.status} · ${b.loaded} filas · ${b.cards} tarjetas · ${b.errors} errores · ${b.warnings} avisos · `));
const a=document.createElement('a');a.href=adminBase+`imports/${b.id}/report`;a.textContent='Descargar reporte';li.append(a);
if(b.status!=='active'){const button=document.createElement('button');button.textContent='Activar versión';
button.onclick=async()=>{if(!confirm(`Revisé el reporte de la versión ${b.id}. ¿Activar ${b.cards} tarjetas?`))return;
const form=new FormData();form.set('reviewed','true');const response=await fetch(adminBase+`imports/${b.id}/activate`,{method:'POST',body:form});
result.textContent=JSON.stringify(await response.json(),null,2);refresh()};li.append(button)}return li}));}
document.querySelector('#upload').onsubmit=async e=>{e.preventDefault();result.textContent='Procesando...';
const r=await fetch(adminBase+'imports',{method:'POST',body:new FormData(e.target)});result.textContent=JSON.stringify(await r.json(),null,2);refresh()};refresh();
</script></html>"""


@app.get("/admin/imports", dependencies=[Depends(admin_auth)])
def list_imports(session=Depends(db)):
    return [batch_data(b) for b in session.scalars(select(Batch).order_by(Batch.id.desc()).limit(50))]


@app.post("/admin/imports", status_code=201, dependencies=[Depends(admin_auth), Depends(same_origin)])
async def upload_import(file: UploadFile = File(...), session=Depends(db)):
    max_bytes = int(os.getenv("MAX_UPLOAD_BYTES", "5242880"))
    content = await file.read(max_bytes + 1)
    if len(content) > max_bytes:
        raise HTTPException(413, "file too large")
    try:
        batch = stage(session, file.filename or "", content)
    except InvalidUpload as exc:
        session.rollback()
        raise HTTPException(422, str(exc)) from exc
    return batch_data(batch)


@app.get("/admin/imports/{batch_id}", dependencies=[Depends(admin_auth)])
def get_import(batch_id: int, session=Depends(db)):
    batch = session.get(Batch, batch_id)
    if not batch:
        raise HTTPException(404, "batch not found")
    return batch_data(batch)


def spreadsheet_text(value):
    s = str(value)
    return "'" + s if s and s[0] in "=+-@\t\r\n" else s


@app.get("/admin/imports/{batch_id}/report", dependencies=[Depends(admin_auth)])
def report(batch_id: int, session=Depends(db)):
    batch = session.get(Batch, batch_id)
    if not batch:
        raise HTTPException(404, "batch not found")
    wb = Workbook()
    summary = wb.active
    summary.title = "SUMMARY"
    for key, value in batch_data(batch).items():
        summary.append([key, spreadsheet_text(value)])
    loaded = wb.create_sheet("LOADED")
    loaded.append(["source_row", "EID", "full_name", "facility_code", "card_number", "source"])
    for c in session.scalars(select(Credential).where(Credential.batch_id == batch_id).order_by(Credential.source_row)):
        loaded.append([c.source_row, spreadsheet_text(c.eid), spreadsheet_text(c.full_name), c.facility_code, c.card_number, c.source])
    sheets = {level: wb.create_sheet(level) for level in ("ERRORS", "WARNINGS", "DUPLICATES")}
    for sheet in sheets.values():
        sheet.append(["source_row", "code", "detail"])
    for issue in session.scalars(select(Issue).where(Issue.batch_id == batch_id).order_by(Issue.source_row)):
        category = "DUPLICATES" if "DUPLICATE" in issue.code or issue.code == "AMBIGUOUS_CARD" else issue.level.upper() + "S"
        sheets[category].append([issue.source_row, issue.code, spreadsheet_text(issue.detail)])
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="import-{batch_id}-report.xlsx"',
                                      "Cache-Control": "no-store"})


@app.post("/admin/imports/{batch_id}/activate", dependencies=[Depends(admin_auth), Depends(same_origin)])
def activate(batch_id: int, reviewed: bool = Form(False), session=Depends(db)):
    if not reviewed:
        raise HTTPException(422, "review the report before activation")
    batch = session.get(Batch, batch_id)
    if not batch:
        raise HTTPException(404, "batch not found")
    current = session.get(Active, 1)
    if current:
        previous = session.get(Batch, current.batch_id)
        if previous and batch.cards_count < previous.cards_count * 0.7:
            raise HTTPException(409, "dataset has over 30% fewer cards than active version")
        if previous:
            previous.status = "inactive"
        current.batch_id = batch.id
    else:
        session.add(Active(id=1, batch_id=batch.id))
    batch.status = "active"
    session.commit()
    return batch_data(batch)
